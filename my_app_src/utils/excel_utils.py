"""Excel 工具（原独立工具 excel.py 迁入）"""

import re
from datetime import datetime


def col_idx_to_letter(idx):
    letter = ""
    while idx > 0:
        idx, remainder = divmod(idx - 1, 26)
        letter = chr(65 + remainder) + letter
    return letter


def to_float_safe(v):
    if v is None:
        return None
    try:
        return float(v)
    except Exception:
        return None


def find_row_by_keyword(ws, col_idx, keyword, max_row=500):
    for row in range(1, max_row + 1):
        val = ws.cell(row=row, column=col_idx).value
        if val is not None and keyword in str(val):
            return row
    return None


def find_price_col(ws, date_obj, max_col=1000):
    for col in range(1, max_col):
        val = ws.cell(row=3, column=col).value
        if val is None:
            continue
        cell_date = None
        if isinstance(val, datetime):
            cell_date = val.date()
        elif isinstance(val, (int, float)):
            try:
                from openpyxl.utils.datetime import from_excel

                cell_date = from_excel(val).date()
            except Exception:
                from datetime import timedelta

                base = datetime(1899, 12, 30).date()
                cell_date = base + timedelta(days=int(val))
        elif isinstance(val, str):
            m = re.search(r"(\d{4})[-/年](\d{1,2})[-/月](\d{1,2})", val)
            if m:
                try:
                    cell_date = datetime(
                        int(m.group(1)), int(m.group(2)), int(m.group(3))
                    ).date()
                except Exception:
                    pass
        if cell_date == date_obj:
            return col
    return None


def find_sheet_by_keywords(wb, keywords, exact_names=None):
    """
    在 wb 中查找匹配的工作表对象。
    - 先按 exact_names 精确匹配
    - 再按 keywords 模糊匹配（sheet 名须包含全部关键字）

    keywords: List[str]，全部包含才命中（如 ["模拟结算", "汇总"]）
    exact_names: List[str]，按顺序尝试
    返回 Worksheet 或 None
    """
    names = wb.sheetnames

    if exact_names:
        for n in exact_names:
            if n in names:
                return wb[n]

    if keywords:
        for n in names:
            if all(k in n for k in keywords):
                return wb[n]

    return None
