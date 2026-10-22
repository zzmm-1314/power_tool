"""现代模拟结算单更新（原 sim_v2.py）"""

import os
import re
import openpyxl
from datetime import datetime
from utils.dates import parse_sheet_name_date, get_month_days, get_tenday_value
from utils.excel_utils import find_price_col, to_float_safe, find_sheet_by_keywords
from utils.number_format import apply_number_format_v2


def fill_simulation_data_v2(
    spot_file,
    stat_file,
    strategy_file,
    info_spot_file,
    info_contract_file,
    target_file,
    annual_price,
    monthly_price,
    tenday_prices,
    selected_dates=None,
    log_func=print,
    progress_func=None,
):
    for fp in [
        spot_file,
        stat_file,
        strategy_file,
        info_spot_file,
        info_contract_file,
        target_file,
    ]:
        if not os.path.isfile(fp):
            log_func(f"错误：文件不存在 {fp}")
            return

    try:
        wb_spot = openpyxl.load_workbook(spot_file, data_only=True)
    except Exception as e:
        log_func(f"读取现货申报表失败: {e}")
        return
    spot_map = {}
    for name in wb_spot.sheetnames:
        md = parse_sheet_name_date(name)
        if md:
            spot_map[md] = name
    if not spot_map:
        log_func("错误：现货申报表无日期工作表")
        wb_spot.close()
        return

    try:
        wb_stat = openpyxl.load_workbook(stat_file, data_only=True, read_only=True)
        ws_stat = wb_stat["中长期持仓详细"]
    except Exception as e:
        log_func(f"读取数据统计表失败: {e}")
        wb_spot.close()
        return

    try:
        wb_strategy = openpyxl.load_workbook(
            strategy_file, data_only=True, read_only=True
        )
        ws_ym = wb_strategy["年月分解表"]
        ws_td = wb_strategy["旬度分解表"]
    except Exception as e:
        log_func(f"读取月度策略失败: {e}")
        wb_spot.close()
        wb_stat.close()
        return

    try:
        wb_info_spot = openpyxl.load_workbook(
            info_spot_file, data_only=True, read_only=True
        )
        ws_spot_price = wb_info_spot["分时价格（现货）"]
    except Exception as e:
        log_func(f"读取日前实时参考表失败: {e}")
        wb_spot.close()
        wb_stat.close()
        wb_strategy.close()
        return

    try:
        wb_info_contract = openpyxl.load_workbook(
            info_contract_file, data_only=True, read_only=True
        )
        ws_contract_price = wb_info_contract["分时价格（现货）"]
    except Exception as e:
        log_func(f"读取合约参考表失败: {e}")
        wb_spot.close()
        wb_stat.close()
        wb_strategy.close()
        wb_info_spot.close()
        return

    try:
        wb_target = openpyxl.load_workbook(target_file)
    except Exception as e:
        log_func(f"加载目标文件失败: {e}")
        wb_spot.close()
        wb_stat.close()
        wb_strategy.close()
        wb_info_spot.close()
        wb_info_contract.close()
        return

    target_sheets = []
    for name in wb_target.sheetnames:
        if not re.match(r"^\d{4}-\d{2}-\d{2}$", name.strip()):
            continue
        try:
            d = datetime.strptime(name.strip(), "%Y-%m-%d").date()
            target_sheets.append((d, name, wb_target[name]))
        except Exception:
            pass
    if not target_sheets:
        log_func("错误：目标文件无日期工作表")
        return
    target_sheets.sort(key=lambda x: x[0])

    if selected_dates:
        sels = set()
        for s in selected_dates:
            try:
                sels.add(datetime.strptime(s, "%Y-%m-%d").date())
            except Exception:
                pass
        target_sheets = [(d, n, ws) for d, n, ws in target_sheets if d in sels]
        if not target_sheets:
            log_func("错误：无匹配日期")
            return

    summary_rows = []
    total = len(target_sheets)
    for loop_idx, (date_obj, date_str, ws_tgt) in enumerate(target_sheets):
        if progress_func:
            progress_func(loop_idx, total)
        log_func(f"处理日期 {date_str}")

        spot_name = spot_map.get((date_obj.month, date_obj.day))
        if not spot_name:
            log_func(f"  现货申报表无 {date_obj.month}月{date_obj.day}日 工作表，跳过")
            continue
        ws_spot = wb_spot[spot_name]

        price_col_spot = find_price_col(ws_spot_price, date_obj)
        price_col_contract = find_price_col(ws_contract_price, date_obj)

        for i in range(48):
            v = to_float_safe(ws_spot.cell(row=3 + i, column=12).value)
            ws_tgt.cell(row=5 + i, column=2, value=v)

        if price_col_spot:
            for i in range(48):
                v = to_float_safe(
                    ws_spot_price.cell(row=5 + i, column=price_col_spot).value
                )
                ws_tgt.cell(row=5 + i, column=3, value=v)

        for i in range(48):
            v = to_float_safe(ws_spot.cell(row=3 + i, column=13).value)
            ws_tgt.cell(row=5 + i, column=5, value=v)

        if price_col_spot:
            for i in range(48):
                v = to_float_safe(
                    ws_spot_price.cell(row=5 + i, column=price_col_spot + 1).value
                )
                ws_tgt.cell(row=5 + i, column=7, value=v)

        days = get_month_days(date_obj)
        base_row = 34 + (days - 28)
        stat_row = base_row + date_obj.day - 1
        for i in range(48):
            v = to_float_safe(ws_stat.cell(row=stat_row, column=4 + i).value)
            ws_tgt.cell(row=5 + i, column=9, value=v)

        for i in range(48):
            ws_tgt.cell(row=5 + i, column=10, value=annual_price)

        if price_col_contract:
            for i in range(48):
                v = to_float_safe(
                    ws_contract_price.cell(row=5 + i, column=price_col_contract).value
                )
                ws_tgt.cell(row=5 + i, column=11, value=v)

        ym_row = 7 + date_obj.day - 1
        for i in range(48):
            v = to_float_safe(ws_ym.cell(row=ym_row, column=6 + i).value)
            ws_tgt.cell(row=5 + i, column=14, value=v)

        for i in range(48):
            ws_tgt.cell(row=5 + i, column=15, value=monthly_price)

        if price_col_contract:
            for i in range(48):
                v = to_float_safe(
                    ws_contract_price.cell(row=5 + i, column=price_col_contract).value
                )
                ws_tgt.cell(row=5 + i, column=16, value=v)

        td_row = 9 + date_obj.day - 1
        for i in range(48):
            v = to_float_safe(ws_td.cell(row=td_row, column=6 + i).value)
            ws_tgt.cell(row=5 + i, column=19, value=v)

        tp = get_tenday_value(date_obj, tenday_prices)
        for i in range(48):
            ws_tgt.cell(row=5 + i, column=20, value=tp)

        if price_col_contract:
            for i in range(48):
                v = to_float_safe(
                    ws_contract_price.cell(row=5 + i, column=price_col_contract).value
                )
                ws_tgt.cell(row=5 + i, column=21, value=v)

        B_sum = 0.0
        D_sum = 0.0
        for i in range(48):
            r = 5 + i
            b = ws_tgt.cell(row=r, column=2).value or 0
            c = ws_tgt.cell(row=r, column=3).value or 0
            e = ws_tgt.cell(row=r, column=5).value or 0
            g = ws_tgt.cell(row=r, column=7).value or 0
            ii = ws_tgt.cell(row=r, column=9).value or 0
            j = ws_tgt.cell(row=r, column=10).value or 0
            k = ws_tgt.cell(row=r, column=11).value or 0
            n = ws_tgt.cell(row=r, column=14).value or 0
            o = ws_tgt.cell(row=r, column=15).value or 0
            p = ws_tgt.cell(row=r, column=16).value or 0
            s = ws_tgt.cell(row=r, column=19).value or 0
            t = ws_tgt.cell(row=r, column=20).value or 0
            u = ws_tgt.cell(row=r, column=21).value or 0

            f = e - b
            d = b * c
            h = f * g
            l = j - k
            m = l * ii
            q = o - p
            rr = q * n
            v = t - u
            w = v * s

            B_sum += b + f
            D_sum += d + h + m + rr + w

        summary_rows.append((date_obj.day, B_sum, D_sum))

        date_label = f"{date_obj.month}月{date_obj.day}日"
        ws_tgt.cell(row=1, column=2, value=date_label)
        ws_tgt.cell(row=53, column=2, value=date_label)

        apply_number_format_v2(ws_tgt)

    if progress_func:
        progress_func(total, total)
    # 汇总 sheet 兼容「模拟结算汇总表」/「模拟结算汇总」/「...结算汇总...」
    ws_summary = find_sheet_by_keywords(
        wb_target,
        keywords=["模拟结算", "汇总"],
        exact_names=["模拟结算汇总表", "模拟结算汇总", "结算汇总", "汇总"],
    )
    if ws_summary is None:
        log_func(
            "加载目标文件失败: 未找到汇总工作表（期望包含「模拟结算」和「汇总」字样），"
            f"现有工作表：{wb_target.sheetnames}"
        )
        wb_target.close()
        wb_spot.close()
        wb_stat.close()
        wb_strategy.close()
        wb_info_spot.close()
        wb_info_contract.close()
        return
    for day, b_sum, d_sum in summary_rows:
        row = 3 + day - 1
        ws_summary.cell(row=row, column=2, value=b_sum)
        ws_summary.cell(row=row, column=4, value=d_sum)

    try:
        wb_target.save(target_file)
        wb_target.close()
        wb_spot.close()
        wb_stat.close()
        wb_strategy.close()
        wb_info_spot.close()
        wb_info_contract.close()
        log_func(f"✅ 完成！共处理 {len(summary_rows)} 天。")
    except Exception as e:
        log_func(f"保存失败: {e}")
