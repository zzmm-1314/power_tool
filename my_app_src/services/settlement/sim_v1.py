"""椒江模拟结算单更新（原 sim_v1.py）"""

import os
import glob
import re
import openpyxl
from datetime import datetime
from utils.dates import (
    extract_date_from_cell_value,
    get_annual_contract_row,
    get_tenday_value,
)
from utils.excel_utils import col_idx_to_letter, find_sheet_by_keywords


def fill_simulation_data(
    detail_folder,
    fixed_power_file,
    strategy_file,
    stat_file,
    info_file,
    target_file,
    annual_price,
    tenday_prices,
    selected_dates=None,
    log_func=print,
    progress_func=None,
):
    if not os.path.isdir(detail_folder):
        log_func("错误：明细文件夹不存在")
        return
    for fp in [fixed_power_file, strategy_file, stat_file, info_file, target_file]:
        if not os.path.isfile(fp):
            log_func(f"错误：文件不存在 {fp}")
            return

    detail_files = glob.glob(os.path.join(detail_folder, "*.xlsx"))
    detail_files = [f for f in detail_files if not os.path.basename(f).startswith("~$")]
    if not detail_files:
        log_func("错误：文件夹无 Excel 文件")
        return

    file_date_pairs = []
    for f in detail_files:
        try:
            wb = openpyxl.load_workbook(f, data_only=True, read_only=True)
            ws = wb.active
            v = ws["A2"].value
            wb.close()
            d = extract_date_from_cell_value(v)
            if d:
                file_date_pairs.append((d, f))
            else:
                log_func(f"警告：{os.path.basename(f)} A2 无法解析日期，跳过")
        except Exception as e:
            log_func(f"读取 {os.path.basename(f)} 失败: {e}")

    if not file_date_pairs:
        log_func("错误：未提取到任何日期")
        return

    file_date_pairs.sort(key=lambda x: x[0])

    if selected_dates:
        sels = set()
        for s in selected_dates:
            try:
                sels.add(datetime.strptime(s, "%Y-%m-%d").date())
            except Exception:
                pass
        file_date_pairs = [(d, f) for d, f in file_date_pairs if d in sels]
        if not file_date_pairs:
            log_func("错误：无匹配日期文件")
            return
        log_func(f"过滤后 {len(file_date_pairs)} 个文件")

    wb_fixed = openpyxl.load_workbook(fixed_power_file, data_only=True, read_only=True)
    ws_fixed = wb_fixed["数据"]
    wb_strategy = openpyxl.load_workbook(strategy_file, data_only=True, read_only=True)
    ws_strategy = wb_strategy["旬度分解表"]
    wb_stat = openpyxl.load_workbook(stat_file, data_only=True, read_only=True)
    ws_stat = wb_stat["中长期持仓详细"]
    wb_info = openpyxl.load_workbook(info_file, data_only=True, read_only=True)
    ws_info = wb_info["分时价格（现货）"]

    wb_target = openpyxl.load_workbook(target_file)
    ws_sim = find_sheet_by_keywords(
        wb_target,
        keywords=["模拟结算"],
        exact_names=["模拟结算", "模拟结算明细"],
    )
    ws_summary = find_sheet_by_keywords(
        wb_target,
        keywords=["模拟结算", "汇总"],
        exact_names=["模拟结算汇总表", "模拟结算汇总"],
    )
    if ws_sim is None or ws_summary is None:
        log_func(
            "加载目标文件失败: 缺少「模拟结算」或「模拟结算汇总」工作表，"
            f"现有工作表：{wb_target.sheetnames}"
        )
        wb_target.close()
        wb_fixed.close()
        wb_strategy.close()
        wb_stat.close()
        wb_info.close()
        return
    COL_OFFSETS = {
        "T": 0,
        "U": 1,
        "W": 3,
        "Y": 5,
        "AA": 7,
        "AB": 8,
        "AC": 9,
        "AF": 12,
        "AG": 13,
        "AH": 14,
    }
    START_COL = 20
    COL_OFFSET_PER_DAY = 18
    summary_rows = []

    total = len(file_date_pairs)
    for loop_idx, (date_obj, file_path) in enumerate(file_date_pairs):
        if progress_func:
            progress_func(loop_idx, total)
        log_func(f"处理日期 {date_obj}：{os.path.basename(file_path)}")

        day_index = date_obj.day - 1
        base_col = START_COL + day_index * COL_OFFSET_PER_DAY

        col_T = base_col + COL_OFFSETS["T"]
        col_U = base_col + COL_OFFSETS["U"]
        col_W = base_col + COL_OFFSETS["W"]
        col_Y = base_col + COL_OFFSETS["Y"]
        col_AA = base_col + COL_OFFSETS["AA"]
        col_AB = base_col + COL_OFFSETS["AB"]
        col_AC = base_col + COL_OFFSETS["AC"]
        col_AF = base_col + COL_OFFSETS["AF"]
        col_AG = base_col + COL_OFFSETS["AG"]
        col_AH = base_col + COL_OFFSETS["AH"]

        try:
            wb_detail = openpyxl.load_workbook(
                file_path, data_only=True, read_only=True
            )
            ws_detail = wb_detail.active
        except Exception as e:
            log_func(f"  读取失败，跳过: {e}")
            continue

        pwr_idx = 2 + day_index * 2
        pwr_letter = col_idx_to_letter(pwr_idx)
        power_values = []
        for row in range(3, 99):
            v = ws_fixed[f"{pwr_letter}{row}"].value
            power_values.append(float(v) if v is not None else 0.0)
        energy_48 = [
            (power_values[i] + power_values[i + 1]) * 0.25 for i in range(0, 96, 2)
        ]

        detail_U = [ws_detail.cell(row=6 + i, column=7).value for i in range(48)]
        detail_W = [ws_detail.cell(row=54 + i, column=5).value for i in range(48)]
        detail_Y = [ws_detail.cell(row=54 + i, column=7).value for i in range(48)]

        row_annual = get_annual_contract_row(date_obj)
        annual_values = []
        for ci in range(4, 52):
            v = ws_stat.cell(row=row_annual, column=ci).value
            annual_values.append(float(v) if v is not None else 0.0)

        tenday_row = 9 + day_index
        tenday_values = []
        for ci in range(6, 54):
            v = ws_strategy.cell(row=tenday_row, column=ci).value
            tenday_values.append(float(v) if v is not None else 0.0)

        date_col = None
        for ci in range(3, 1000, 3):
            val = ws_info.cell(row=3, column=ci).value
            if val is None:
                continue
            cd = None
            if isinstance(val, datetime):
                cd = val.date()
            elif isinstance(val, (int, float)):
                try:
                    from openpyxl.utils.datetime import from_excel

                    cd = from_excel(val).date()
                except Exception:
                    from datetime import timedelta

                    cd = datetime(1899, 12, 30).date() + timedelta(days=int(val))
            elif isinstance(val, str):
                m = re.search(r"(\d{4})[-/年](\d{1,2})[-/月](\d{1,2})", val)
                if m:
                    try:
                        cd = datetime(
                            int(m.group(1)), int(m.group(2)), int(m.group(3))
                        ).date()
                    except Exception:
                        pass
            if cd == date_obj:
                date_col = ci
                break

        avg_prices = None
        if date_col:
            avg_prices = [
                ws_info.cell(row=r, column=date_col).value for r in range(5, 53)
            ]
            avg_prices = [float(v) if v is not None else 0.0 for v in avg_prices]

        sum_w = 0.0
        for i in range(48):
            r = 5 + i
            val_t = energy_48[i]
            ws_sim.cell(row=r, column=col_T, value=val_t)

            val_u = detail_U[i] if detail_U[i] is not None else 0.0
            ws_sim.cell(row=r, column=col_U, value=val_u)

            val_w = detail_W[i] if detail_W[i] is not None else 0.0
            ws_sim.cell(row=r, column=col_W, value=val_w)
            sum_w += val_w

            val_y = detail_Y[i] if detail_Y[i] is not None else 0.0
            ws_sim.cell(row=r, column=col_Y, value=val_y)

            ws_sim.cell(row=r, column=col_AA, value=annual_values[i])
            ws_sim.cell(row=r, column=col_AB, value=annual_price)

            if avg_prices is not None:
                ws_sim.cell(row=r, column=col_AC, value=avg_prices[i])
                ws_sim.cell(row=r, column=col_AH, value=avg_prices[i])

            ws_sim.cell(row=r, column=col_AF, value=tenday_values[i])
            tp = get_tenday_value(date_obj, tenday_prices)
            ws_sim.cell(row=r, column=col_AG, value=tp)

        u58_val = sum_w
        sum_v = sum_z = sum_ae = sum_aj = 0.0
        for r in range(5, 53):
            t = ws_sim.cell(row=r, column=col_T).value or 0.0
            u = ws_sim.cell(row=r, column=col_U).value or 0.0
            w = ws_sim.cell(row=r, column=col_W).value or 0.0
            y = ws_sim.cell(row=r, column=col_Y).value or 0.0
            aa = ws_sim.cell(row=r, column=col_AA).value or 0.0
            ab = ws_sim.cell(row=r, column=col_AB).value or 0.0
            ac = ws_sim.cell(row=r, column=col_AC).value or 0.0
            af = ws_sim.cell(row=r, column=col_AF).value or 0.0
            ag = ws_sim.cell(row=r, column=col_AG).value or 0.0
            ah = ws_sim.cell(row=r, column=col_AH).value or 0.0

            x = w - t
            v = t * u
            z = y * x
            ad = ab - ac
            ae = aa * ad
            ai = ag - ah
            aj = af * ai

            sum_v += v
            sum_z += z
            sum_ae += ae
            sum_aj += aj

        ae58_val = sum_v + sum_z + sum_ae + sum_aj
        summary_rows.append((date_obj.day, u58_val, ae58_val))
        wb_detail.close()

    if progress_func:
        progress_func(total, total)

    for day, u_val, ae_val in summary_rows:
        row = 3 + day - 1
        ws_summary.cell(row=row, column=2, value=u_val)
        ws_summary.cell(row=row, column=4, value=ae_val)

    try:
        wb_target.save(target_file)
        wb_target.close()
        log_func(f"✅ 完成！共处理 {len(summary_rows)} 天。")
    except Exception as e:
        log_func(f"保存失败: {e}")

    wb_fixed.close()
    wb_strategy.close()
    wb_stat.close()
    wb_info.close()
