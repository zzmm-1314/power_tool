"""实际结算单更新（原 task1.py）"""

import os
import glob
import openpyxl
from utils.dates import normalize_date, extract_date_from_cell_value
from utils.excel_utils import find_row_by_keyword


def fill_data(
    folder_path,
    target_file,
    selected_dates=None,
    log_func=print,
    progress_func=None,
):
    if not os.path.isdir(folder_path):
        log_func(f"错误：文件夹不存在 {folder_path}")
        return
    if not os.path.isfile(target_file):
        log_func(f"错误：目标文件不存在 {target_file}")
        return

    all_files = glob.glob(os.path.join(folder_path, "*.xlsx"))
    all_files = [f for f in all_files if not os.path.basename(f).startswith("~$")]
    all_files.sort(key=lambda f: os.path.basename(f))
    if not all_files:
        log_func("错误：文件夹中没有有效的 .xlsx 文件")
        return
    log_func(f"共找到 {len(all_files)} 个 Excel 文件")

    log_func("开始提取 B4 日期...")
    file_info = []
    for f in all_files:
        try:
            wb = openpyxl.load_workbook(f, data_only=True, read_only=True)
            ws = wb.active
            b4 = ws["B4"].value
            wb.close()
            d = extract_date_from_cell_value(b4)
            if d:
                file_info.append((f, d.day, d.strftime("%Y-%m-%d")))
                log_func(f"文件 {os.path.basename(f)} B4 -> {d}")
            else:
                log_func(f"警告：{os.path.basename(f)} B4 无法解析日期，跳过")
        except Exception as e:
            log_func(f"读取 {os.path.basename(f)} 日期失败: {e}，跳过")

    if not file_info:
        log_func("错误：未提取到任何日期")
        return

    if selected_dates:
        norm = [normalize_date(d) for d in selected_dates]
        file_info = [(f, d, ds) for f, d, ds in file_info if ds in norm]
        if not file_info:
            log_func("错误：没有匹配指定日期的文件")
            return
        log_func(f"按日期过滤后，剩余 {len(file_info)} 个文件")

    try:
        wb_tgt = openpyxl.load_workbook(target_file)
        ws_info = wb_tgt["信息汇总（现货每日结算表）"]
        ws_settle = wb_tgt["结算明细表"]
    except Exception as e:
        log_func(f"加载目标文件失败: {e}")
        return

    total = len(file_info)
    for loop_idx, (path, day, date_str) in enumerate(file_info):
        if progress_func:
            progress_func(loop_idx, total)
        log_func(f"处理 {os.path.basename(path)} (日期 {date_str}, 日 {day})")

        try:
            wb_src = openpyxl.load_workbook(path, data_only=True)
        except Exception as e:
            log_func(f"  读取失败，跳过: {e}")
            continue
        ws_src = wb_src.active

        info_row = 3 + day - 1
        settle_row = 4 + day - 1

        val = ws_src["C8"].value
        if val is not None:
            ws_info[f"B{info_row}"] = val

        r = find_row_by_keyword(ws_src, 2, "总计")
        if r:
            v = ws_src.cell(row=r, column=2 + 7).value
            if v is not None:
                ws_info[f"D{info_row}"] = v

        r = find_row_by_keyword(ws_src, 3, "日前市场")
        if r:
            v = ws_src.cell(row=r, column=3 + 2).value
            if v is not None:
                ws_settle[f"C{settle_row}"] = v
            v = ws_src.cell(row=r, column=3 + 4).value
            if v is not None:
                ws_settle[f"E{settle_row}"] = v

        r = find_row_by_keyword(ws_src, 3, "实时市场")
        if r:
            v = ws_src.cell(row=r, column=3 + 4).value
            if v is not None:
                ws_settle[f"G{settle_row}"] = v

        r = find_row_by_keyword(ws_src, 4, "年度集中合约")
        if r:
            for off, col in [(1, "M"), (3, "N"), (5, "O")]:
                v = ws_src.cell(row=r, column=4 + off).value
                if v is not None:
                    ws_settle[f"{col}{settle_row}"] = v

        r = find_row_by_keyword(ws_src, 4, "月度集中合约")
        if r:
            for off, col in [(1, "P"), (3, "Q"), (5, "R")]:
                v = ws_src.cell(row=r, column=4 + off).value
                if v is not None:
                    ws_settle[f"{col}{settle_row}"] = v

        r = find_row_by_keyword(ws_src, 4, "日内滚撮合约")
        if r:
            for off, col in [(1, "S"), (3, "T"), (5, "U")]:
                v = ws_src.cell(row=r, column=4 + off).value
                if v is not None:
                    ws_settle[f"{col}{settle_row}"] = v

        wb_src.close()

    if progress_func:
        progress_func(total, total)

    try:
        wb_tgt.save(target_file)
        wb_tgt.close()
        log_func(f"✅ 完成！共处理 {len(file_info)} 天数据。")
    except Exception as e:
        log_func(f"保存失败: {e}")
