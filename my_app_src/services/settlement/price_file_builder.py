"""从数据目录构建结算单所需的「分时价格（现货）」Excel
输出格式（兼容 sim_v1 / sim_v2 的读取逻辑）：
- 第 3 行：日期表头（col 3, 6, 9...，每日期占 3 列）
- 第 5-52 行：日前价格在日期列本身
- 第 5-52 行：实时价格在日期列 + 1 列（仅 with_real_time）
"""

import os
import time
import uuid
from datetime import datetime, timedelta
from typing import Dict, List, Optional

import openpyxl
from openpyxl.styles import Font

from services.day_ahead_parser import (
    _norm_time,
    _time_to_slot_index,
    parse_unified_price_excel,
)
from services.day_ahead_service import TempManager, read_node_excel

# ============================================================
#                    基础工具
# ============================================================


def _date_list(start_date: str, end_date: str, max_days: int = 60) -> List[str]:
    try:
        s = datetime.strptime(start_date, "%Y-%m-%d").date()
        e = datetime.strptime(end_date, "%Y-%m-%d").date()
    except Exception:
        return []
    if e < s:
        s, e = e, s
    result = []
    cur = s
    while cur <= e and len(result) < max_days:
        result.append(cur.strftime("%Y-%m-%d"))
        cur += timedelta(days=1)
    return result


def _raw_to_48(raw: List[Dict]) -> List[Optional[float]]:
    """raw [{time, value}] → 48 槽位数组"""
    slots: List[Optional[float]] = [None] * 48
    for pt in raw or []:
        t = _norm_time(pt.get("time"))
        if not t:
            continue
        idx = _time_to_slot_index(t)
        if idx is None:
            continue
        v = pt.get("value")
        slots[idx] = float(v) if v is not None else None
    return slots


def _workspace_dir() -> str:
    d = os.path.join(TempManager.base_dir(), "_settlement_workspace")
    os.makedirs(d, exist_ok=True)
    return d


def _cleanup_old_files(workspace: str, max_age_hours: int = 24) -> None:
    """清理超过指定时长的旧临时文件"""
    now = time.time()
    try:
        for fn in os.listdir(workspace):
            fp = os.path.join(workspace, fn)
            try:
                if now - os.path.getmtime(fp) > max_age_hours * 3600:
                    os.remove(fp)
            except OSError:
                pass
    except OSError:
        pass


# ============================================================
#                    数据加载
# ============================================================


def _load_unified(dates: List[str]) -> Dict[str, List[Optional[float]]]:
    """从 temp/day-ahead/unified 读日前统一结算点
    返回 {date: [48 点]}
    """
    base = os.path.join(TempManager.base_dir(), "day-ahead", "unified")
    if not os.path.isdir(base):
        return {}

    date_set = set(dates)
    result: Dict[str, List[Optional[float]]] = {}

    for fn in sorted(os.listdir(base)):
        if not fn.lower().endswith((".xlsx", ".xls")):
            continue
        fp = os.path.join(base, fn)
        try:
            items = parse_unified_price_excel(fp)
        except Exception as e:
            print(f"[PriceBuilder] 解析 {fn} 失败: {e}")
            continue
        for it in items:
            if it["date"] in date_set and it["date"] not in result:
                result[it["date"]] = _raw_to_48(it["raw"])

    return result


def _load_node_pair(
    node_name: str, dates: List[str]
) -> Dict[str, Dict[str, List[Optional[float]]]]:
    """读取同一节点的日前与实时
    返回 {date: {"dayAhead": [...], "realTime": [...]}}
    """
    date_set = set(dates)
    result: Dict[str, Dict] = {}

    # 日前节点
    try:
        day_items = read_node_excel(node_name, base_subdir="day-ahead/plant")
    except Exception:
        day_items = []
    for it in day_items:
        if it["date"] in date_set:
            result.setdefault(it["date"], {})["dayAhead"] = _raw_to_48(it["raw"])

    # 实时节点
    try:
        rt_items = read_node_excel(node_name, base_subdir="spread/plant")
    except Exception:
        rt_items = []
    for it in rt_items:
        if it["date"] in date_set:
            result.setdefault(it["date"], {})["realTime"] = _raw_to_48(it["raw"])

    return result


# ============================================================
#                    写 Excel
# ============================================================


def _write_price_sheet(
    data_by_date: Dict, out_path: str, with_real_time: bool = False
) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "分时价格（现货）"

    sorted_dates = sorted(data_by_date.keys())

    for i, date_str in enumerate(sorted_dates):
        col = 3 + i * 3  # 日期起始列
        try:
            date_obj = datetime.strptime(date_str, "%Y-%m-%d").date()
        except Exception:
            continue

        c = ws.cell(row=3, column=col, value=date_obj)
        c.number_format = "yyyy-mm-dd"
        c.font = Font(bold=True)

        data = data_by_date[date_str]

        if with_real_time:
            day_slots = data.get("dayAhead") or [None] * 48
            rt_slots = data.get("realTime") or [None] * 48
            for j in range(48):
                if j < len(day_slots) and day_slots[j] is not None:
                    ws.cell(row=5 + j, column=col, value=day_slots[j])
                if j < len(rt_slots) and rt_slots[j] is not None:
                    ws.cell(row=5 + j, column=col + 1, value=rt_slots[j])
        else:
            slots = data if isinstance(data, list) else [None] * 48
            for j in range(48):
                if j < len(slots) and slots[j] is not None:
                    ws.cell(row=5 + j, column=col, value=slots[j])

    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    wb.save(out_path)
    wb.close()


# ============================================================
#                    对外接口
# ============================================================


def build_unified_file(start_date: str, end_date: str) -> Dict:
    """日前统一结算点 → 「分时价格（现货）」Excel
    用于椒江 info_file / 现代 info_contract_file
    """
    dates = _date_list(start_date, end_date)
    if not dates:
        return {"ok": False, "error": "日期范围无效"}

    data = _load_unified(dates)
    if not data:
        return {
            "ok": False,
            "error": "数据目录「day-ahead/unified」中无日前统一结算点数据",
        }

    workspace = _workspace_dir()
    _cleanup_old_files(workspace)

    fname = f"分时价格_统一结算点_{uuid.uuid4().hex[:8]}.xlsx"
    out_path = os.path.join(workspace, fname)
    _write_price_sheet(data, out_path, with_real_time=False)

    return {"ok": True, "path": out_path, "dateCount": len(data)}


def build_plant_pair_file(node_name: str, start_date: str, end_date: str) -> Dict:
    """节点（日前 + 实时） → 「分时价格（现货）」Excel
    用于现代 info_spot_file
    """
    if not node_name:
        return {"ok": False, "error": "缺少节点名"}

    dates = _date_list(start_date, end_date)
    if not dates:
        return {"ok": False, "error": "日期范围无效"}

    data = _load_node_pair(node_name, dates)
    if not data:
        return {
            "ok": False,
            "error": f"节点「{node_name}」在范围内无日前/实时数据",
        }

    workspace = _workspace_dir()
    _cleanup_old_files(workspace)

    safe_name = node_name.replace("/", "_").replace("\\", "_").replace(":", "_")
    fname = f"分时价格_{safe_name}_{uuid.uuid4().hex[:8]}.xlsx"
    out_path = os.path.join(workspace, fname)
    _write_price_sheet(data, out_path, with_real_time=True)

    return {
        "ok": True,
        "path": out_path,
        "dateCount": len(data),
        "nodeName": node_name,
    }
