"""滚撮分析：从 HTML 片段 + 日前统一结算点均价做价差分析"""

import io
import re
from typing import Dict, List, Optional, Tuple

import pandas as pd
from bs4 import BeautifulSoup

from services.day_ahead_service import _scan_unified_temp_files

COLUMNS = [
    "时段类型",
    "最新价",
    "买3价",
    "买3量",
    "买2价",
    "买2量",
    "买1价",
    "买1量",
    "卖1价",
    "卖1量",
    "卖2价",
    "卖2量",
    "卖3价",
    "卖3量",
    "未成交量(日均)",
    "成交方向",
    "可买量(日均)",
    "可卖量(日均)",
    "可申报笔数",
]


# ==================== HTML 解析 ====================


def parse_html_to_df(html: str) -> pd.DataFrame:
    soup = BeautifulSoup(html, "lxml")
    body_wrapper = soup.find("div", class_="el-table__body-wrapper")
    table = (
        body_wrapper.find("table")
        if body_wrapper
        else soup.find("table", class_="el-table__body")
    )
    if not table:
        raise ValueError("未找到数据表格，请检查 HTML 片段是否完整。")
    tbody = table.find("tbody")
    if not tbody:
        raise ValueError("表格中没有 <tbody> 部分。")
    rows = tbody.find_all("tr", class_="el-table__row")
    if not rows:
        raise ValueError("表格中没有数据行。")

    data = []
    for row in rows:
        tds = row.find_all("td")
        row_data = []
        for idx, td in enumerate(tds[:13]):
            raw = td.get_text(separator="\n", strip=True)
            parts = [p.strip() for p in re.split(r"[\n\r]+", raw) if p.strip()]
            if 2 <= idx <= 7:
                if len(parts) >= 2:
                    row_data.extend(parts[:2])
                elif len(parts) == 1:
                    row_data.extend([parts[0], ""])
                else:
                    row_data.extend(["", ""])
            else:
                row_data.append(parts[0] if parts else "")
        while len(row_data) < len(COLUMNS):
            row_data.append("")
        data.append(row_data[: len(COLUMNS)])

    return pd.DataFrame(data, columns=COLUMNS)


def parse_hour_label(label) -> Optional[int]:
    if not label:
        return None
    m = re.match(r"(\d{1,2})", str(label))
    return int(m.group(1)) if m else None


# ==================== 日前均价（从数据目录） ====================


def _load_unified_24h_by_date() -> Dict[str, List[Optional[float]]]:
    """{ '2026-10-20': [24 点价格] }"""
    items = _scan_unified_temp_files()
    result: Dict[str, List[Optional[float]]] = {}
    for it in items:
        d = it["date"]
        hourly = it.get("hourly") or []
        vals: List[Optional[float]] = []
        for h in hourly[:24]:
            v = h.get("value")
            vals.append(float(v) if v is not None else None)
        # 补齐到 24 个
        while len(vals) < 24:
            vals.append(None)
        result[d] = vals
    return result


def get_avg_prices(
    days: int, ref_date: Optional[str] = None
) -> Tuple[List[str], List[Optional[float]]]:
    """返回 (选中的日期列表, 24 点均值)"""
    data = _load_unified_24h_by_date()
    if not data:
        raise ValueError(
            "数据目录中没有「日前统一结算点」数据，"
            "请先在「日前电价 / 统一结算点」上传。"
        )

    dates = sorted(data.keys())
    if ref_date:
        dates = [d for d in dates if d <= ref_date]
    if not dates:
        raise ValueError(f"参考日期 {ref_date} 之前没有可用数据")

    selected = dates[-days:] if days > 0 else dates

    avg: List[Optional[float]] = []
    for h in range(24):
        vals = [data[d][h] for d in selected if data[d][h] is not None]
        avg.append(sum(vals) / len(vals) if vals else None)
    return selected, avg


# ==================== 主分析 ====================


def _f(v):
    try:
        if v is None or (isinstance(v, float) and pd.isna(v)):
            return None
        return float(v)
    except (TypeError, ValueError):
        return None


def analyze(html: str, avg_days: int = 3, ref_date: Optional[str] = None) -> Dict:
    df = parse_html_to_df(html)
    selected_dates, avg_prices = get_avg_prices(avg_days, ref_date)

    rows = []
    for _, row in df.iterrows():
        h = parse_hour_label(row["时段类型"])
        if h is None or h >= 24:
            continue

        bid1 = _f(row.get("买1价"))
        ask1 = _f(row.get("卖1价"))
        avg = avg_prices[h] if h < len(avg_prices) else None

        rows.append(
            {
                "hour": h,
                "hour_label": row["时段类型"],
                "latest": _f(row.get("最新价")),
                "bid3": _f(row.get("买3价")),
                "bid3_vol": _f(row.get("买3量")),
                "bid2": _f(row.get("买2价")),
                "bid2_vol": _f(row.get("买2量")),
                "bid1": bid1,
                "bid1_vol": _f(row.get("买1量")),
                "ask1": ask1,
                "ask1_vol": _f(row.get("卖1量")),
                "ask2": _f(row.get("卖2价")),
                "ask2_vol": _f(row.get("卖2量")),
                "ask3": _f(row.get("卖3价")),
                "ask3_vol": _f(row.get("卖3量")),
                "unfilled": _f(row.get("未成交量(日均)")),
                "direction": row.get("成交方向", ""),
                "buyable": _f(row.get("可买量(日均)")),
                "sellable": _f(row.get("可卖量(日均)")),
                "orderable": _f(row.get("可申报笔数")),
                "avg_price": avg,
                "bid1_diff": (
                    round(bid1 - avg, 3)
                    if (bid1 is not None and avg is not None)
                    else None
                ),
                "ask1_diff": (
                    round(ask1 - avg, 3)
                    if (ask1 is not None and avg is not None)
                    else None
                ),
            }
        )
    rows.sort(key=lambda x: x["hour"])

    return {
        "dates_used": selected_dates,
        "avg_days": avg_days,
        "rows": rows,
    }


# ==================== Excel 导出 ====================


def _hex_color_for_excel(value, min_val, max_val) -> Optional[str]:
    if value is None or min_val is None or max_val is None:
        return None
    try:
        value = float(value)
        min_val = float(min_val)
        max_val = float(max_val)
    except (ValueError, TypeError):
        return None
    if max_val == min_val:
        return "FFEB84"
    ratio = max(0.0, min(1.0, (value - min_val) / (max_val - min_val)))
    if ratio < 0.5:
        t = ratio * 2
        r = int(0x63 + (0xFF - 0x63) * t)
        g = int(0xBE + (0xEB - 0xBE) * t)
        b = int(0x7B + (0x84 - 0x7B) * t)
    else:
        t = (ratio - 0.5) * 2
        r = int(0xFF + (0xF8 - 0xFF) * t)
        g = int(0xEB + (0x69 - 0xEB) * t)
        b = int(0x84 + (0x6B - 0x84) * t)
    return f"{r:02X}{g:02X}{b:02X}"


def export_excel(
    html: str, avg_days: int = 3, ref_date: Optional[str] = None
) -> io.BytesIO:
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

    res = analyze(html, avg_days, ref_date)
    df = pd.DataFrame(res["rows"])
    df = df.rename(
        columns={
            "hour_label": "时段类型",
            "latest": "最新价",
            "bid3": "买3价",
            "bid3_vol": "买3量",
            "bid2": "买2价",
            "bid2_vol": "买2量",
            "bid1": "买1价",
            "bid1_vol": "买1量",
            "ask1": "卖1价",
            "ask1_vol": "卖1量",
            "ask2": "卖2价",
            "ask2_vol": "卖2量",
            "ask3": "卖3价",
            "ask3_vol": "卖3量",
            "unfilled": "未成交量(日均)",
            "direction": "成交方向",
            "buyable": "可买量(日均)",
            "sellable": "可卖量(日均)",
            "orderable": "可申报笔数",
            "avg_price": "日前均价",
            "bid1_diff": "买1价-日前均价",
            "ask1_diff": "卖1价-日前均价",
        }
    )
    df = df.drop(columns=["hour"], errors="ignore")
    cols = list(df.columns)

    wb = Workbook()
    ws = wb.active
    ws.title = "滚撮分析"

    header_fill = PatternFill(
        start_color="EEF2FF", end_color="EEF2FF", fill_type="solid"
    )
    header_font = Font(bold=True, color="1E293B", size=11)
    thin = Side(border_style="thin", color="E8ECF3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")

    for ci, c in enumerate(cols, 1):
        cell = ws.cell(row=1, column=ci, value=c)
        cell.font = header_font
        cell.alignment = center
        cell.fill = header_fill
        cell.border = border

    for ri, row in df.iterrows():
        for ci, c in enumerate(cols, 1):
            v = row[c]
            if v is None or (isinstance(v, float) and pd.isna(v)):
                v = None
            cell = ws.cell(row=ri + 2, column=ci, value=v)
            cell.alignment = center
            cell.border = border

    # 色阶
    highlight_cols = ["买1价", "卖1价", "买1价-日前均价", "卖1价-日前均价"]
    for hc in highlight_cols:
        if hc not in cols:
            continue
        ci = cols.index(hc) + 1
        vals = [
            float(v)
            for v in df[hc]
            if v is not None and not (isinstance(v, float) and pd.isna(v))
        ]
        if not vals:
            continue
        min_v, max_v = min(vals), max(vals)
        for ri in range(2, len(df) + 2):
            cell = ws.cell(row=ri, column=ci)
            hex_color = _hex_color_for_excel(cell.value, min_v, max_v)
            if hex_color:
                cell.fill = PatternFill(
                    start_color=hex_color, end_color=hex_color, fill_type="solid"
                )

    # 列宽
    for ci, c in enumerate(cols, 1):
        max_len = len(str(c)) * 2
        for ri in range(2, len(df) + 2):
            v = ws.cell(row=ri, column=ci).value
            if v is not None:
                s = str(v)
                w = sum(2 if ord(ch) > 127 else 1 for ch in s)
                max_len = max(max_len, w)
        letter = ws.cell(row=1, column=ci).column_letter
        ws.column_dimensions[letter].width = min(max_len + 3, 22)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out
